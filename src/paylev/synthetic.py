"""Fake Damodaran files with known effects, for the tests (nothing here is used for real results).

`write_fake_damodaran` writes the five datasets for every year in the same three layouts as the real archive
(1999-2012 plain `Sheet1`; 2013-2018 `Sheet1` with a "Date updated" stamp; 2019+ `Variables & FAQ` +
`Industry Averages`), with the real column names of each era, so the loader is exercised the way the real files
exercise it. The data contain:

- a planted leverage effect and leverage x high-rate interaction in the payout ratio (so the pipeline has to
  find them), with high-rate years that match the committed FRED snapshot and the 3% rule;
- a block of rows in one wacc file shifted against the industry names (like wacc99), which the row-alignment
  check must catch;
- a year whose divfcfe net income is scrambled across industries (like divfcfe07), which the cross-file check
  must catch;
- a financial and a utility industry (excluded), and an industry renamed at the 2013 reclassification
  ("Restaurant" -> "Restaurant/Dining", reconciled by the notebook's rename map);
- a newest file saved without year digits (dated from its stamp, like a fresh download from /datasets/).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl

HIGH_RATE_YEARS = [1999, 2000, 2001, 2005, 2006, 2007, 2023, 2024, 2025]   # fed funds >= 3% in the snapshot
TRUTH = {"lev": 0.3, "lev_x_high": -0.4, "roe": -1.0}                     # planted payout coefficients
SHIFTED = ("wacc", 2004, 3, 9)       # file, year, rows [3, 9) of the industry list shifted by one
SCRAMBLED = ("divfcfe", 2007)        # net income permuted across industries

INDUSTRIES = ["Advertising", "Aerospace/Defense", "Air Transport", "Apparel", "Auto & Truck", "Beverage",
              "Building Materials", "Chemical (Basic)", "Computers/Peripherals", "Electrical Equipment",
              "Entertainment", "Food Processing", "Healthcare Products", "Homebuilding", "Hotel/Gaming",
              "Machinery", "Metals & Mining", "Oil/Gas (Integrated)", "Packaging & Container", "Paper/Forest Products",
              "Retail (General)", "Semiconductor", "Steel", "Trucking", "Restaurant",
              "Bank (Regional)", "Utility (Water)"]
AGGREGATE = "Total Market"


def _era(year):
    return "classic" if year <= 2012 else ("transition" if year <= 2018 else "modern")


def _name(ind, year):
    return "Restaurant/Dining" if ind == "Restaurant" and year >= 2013 else ind


# Column headers per dataset and era: (header text, key into the simulated values). The texts are the real
# ones (odd spacing included), because the loader matches them by keyword.
HEADERS = {
    "divfund": {
        "classic": [("Number of Firms", "n"), ("Dividend Yield", "yield"), ("Dividend Payout", "payout"),
                    (" Market Cap ", "mktcap"), ("ROE", "roe"), ("Insider Holdings", "junk")],
        "transition": [("Number of firms", "n"), ("Dividend Payout", "payout"), ("Dividend Yield", "yield"),
                       ("Market Cap", "mktcap"), ("ROE", "roe"), ("Institutional Holdings", "junk")],
        "modern": [("Number of firms", "n"), ("Total Dividends (US $ millions)", "div"),
                   ("Dividend Payout", "payout"), ("Dividend Yield", "yield"),
                   ("Market Cap (US $ millions)", "mktcap"), ("ROE", "roe")],
    },
    "wacc": {
        "classic": [("Number of Firms", "n"), ("Beta", "junk"), ("E/(D+E)", "eq"), ("Cost of Debt", "junk"),
                    ("Tax Rate", "tax"), ("D/(D+E)", "lev"), ("Cost of Capital", "junk")],
        "transition": [("Number of Firms", "n"), ("Beta", "junk"), ("E/(D+E)", "eq"), ("Tax Rate", "tax"),
                       ("After-tax Cost of Debt", "junk"), ("D/(D+E)", "lev"), ("Cost of Capital", "junk")],
        "modern": [("Number of Firms", "n"), ("Beta", "junk"), ("E/(D+E)", "eq"), ("Tax Rate", "tax"),
                   ("D/(D+E)", "lev"), ("Cost of Capital", "junk"), ("Cost of Capital (Local Currency)", "junk")],
    },
    "divfcfe": {
        "classic": [("Number of firms", "n"), ("  Dividends ", "div"), (" Net Income ", "ni"), ("FCFE", "fcfe"),
                    ("Payout", "payout"), ("Dividends/FCFE", "junk")],
        "transition": [("Number of firms", "n"), ("  Dividends (US $ millions)", "div"),
                       (" Net Income (US $ millions)", "ni"), ("Payout", "payout"),
                       ("Dividends + Buybacks(US $ millions)", "totpay"),
                       ("FCFE (before debt cash flows)(US $ millions)", "fcfe"),
                       ("Net Cash Returned/FCFE (pre-debt)", "junk")],
        "modern": [("Number of firms", "n"), ("  Dividends ", "div"), (" Net Income ", "ni"), ("Payout", "payout"),
                   ("Dividends + Buybacks", "totpay"), ("FCFE (before debt cash flows)", "fcfe"),
                   ("Net Cash Returned/ Net Income", "junk")],
    },
    "capex": {
        "classic": [("Number of Firms", "n"), ("Capital Expenditures", "junk"), ("Depreciation", "junk"),
                    ("Cap Ex/Deprecn", "capex_dep"), ("Net Cap Ex/Sales", "ncx_sales"), ("Sales/Capital", "s_cap")],
        "transition": [("Number of Firms", "n"), ("Capital Expenditures (US $ millions)", "junk"),
                       ("Cap Ex/Deprecn", "capex_dep"), ("Net Cap Ex/Sales", "ncx_sales"),
                       ("Sales/Capital", "s_cap")],
        "modern": [("Number of Firms", "n"), ("Capital Expenditures (US $ millions)", "junk"),
                   ("Cap Ex/Deprecn", "capex_dep"), ("Net Cap Ex/Sales", "ncx_sales"),
                   ("Sales/ Invested Capital (LTM)", "s_cap")],
    },
    "dbtfund": {
        "classic": [("Number of Firms", "n"), ("MV Debt Ratio", "lev_unadj"), ("BV Debt Ratio", "lev_book"),
                    ("Effective Tax Rate", "tax"), ("EBITDA/Value", "ebitda_ev"),
                    ("Fixed Assets/BV of Capital", "ppe"), ("Capital Spending/BV of Capital", "capspend")],
        "transition": [("Number of firms", "n"), ("Book Debt to Capital", "lev_book"),
                       ("Market Debt to Capital (Unadjusted)", "lev_unadj"),
                       ("Market Debt to Capital (adjusted for leases)", "lev"),
                       ("Debt/EBITDA", "de_rep"), ("EBITDA/EV", "ebitda_ev"),
                       ("Net PP&E/Total Assets", "ppe"), ("Capital Spending/Total Assets", "capspend")],
        "modern": [("Number of firms", "n"), ("Book Debt to Capital", "lev_book"),
                   ("Market Debt to Capital (Unadjusted)", "lev_unadj"),
                   ("Market Debt to Capital (adjusted for leases)", "lev"), ("Interest Coverage Ratio", "icr"),
                   ("Debt to EBITDA", "de_rep"), ("EBITDA/EV", "ebitda_ev"), ("Net PP&E/Total Assets", "ppe"),
                   ("Capital Spending/Total Assets", "capspend")],
    },
}


def _simulate(years, seed):
    """Industry x year values, keyed [year][industry][variable]."""
    rng = np.random.default_rng(seed)
    n_ind = len(INDUSTRIES)
    base = {"a": rng.uniform(0.25, 0.5, n_ind), "lev": rng.uniform(0.08, 0.5, n_ind),
            "roe": rng.uniform(0.06, 0.18, n_ind), "n": rng.integers(10, 250, n_ind),
            "ppe": rng.uniform(0.05, 0.7, n_ind), "mktcap": rng.lognormal(10, 1, n_ind)}
    year_fx = {y: rng.normal(0, 0.03) for y in years}
    out = {}
    for y in years:
        high = y in HIGH_RATE_YEARS
        # PP&E is reported on a different basis in each era (the notebook ranks it within the year)
        ppe_scale = {"classic": 1.3, "transition": 1.0, "modern": 0.85}[_era(y)]
        out[y] = {}
        for i, ind in enumerate(INDUSTRIES):
            lev = float(np.clip(base["lev"][i] + rng.normal(0, 0.06), 0.02, 0.85))
            roe = float(np.clip(base["roe"][i] + rng.normal(0, 0.02), 0.02, 0.4))
            payout = (base["a"][i] + year_fx[y] + TRUTH["lev"] * lev + TRUTH["lev_x_high"] * lev * high
                      + TRUTH["roe"] * roe + rng.normal(0, 0.015))
            payout = float(np.clip(payout, 0.02, 1.5))
            mktcap = float(base["mktcap"][i] * np.exp(rng.normal(0, 0.1)))
            ni = mktcap * roe * 0.5
            div = payout * ni
            ppe = float(np.clip(base["ppe"][i] + rng.normal(0, 0.02), 0.01, 0.95)) * ppe_scale
            ebitda_ev = float(rng.uniform(0.06, 0.16))
            lease = 0.03 if 2013 <= y <= 2019 else 0.005     # wacc's D/(D+E) includes leases from 2013 on
            v = {"n": int(base["n"][i] + rng.integers(-3, 4)), "payout": payout, "roe": roe, "mktcap": mktcap,
                 "yield": div / mktcap, "div": div, "ni": ni, "fcfe": ni * rng.uniform(0.6, 1.2),
                 "totpay": div * rng.uniform(1.2, 2.0), "lev": lev, "eq": 1 - lev,
                 "tax": float(rng.uniform(0.1, 0.35)), "lev_unadj": max(lev - lease, 0.0),
                 "lev_book": float(np.clip(lev * 1.6, 0.05, 0.95)), "ebitda_ev": ebitda_ev,
                 "de_rep": lev / ebitda_ev * 1.15, "icr": float(rng.uniform(2, 15)), "ppe": ppe,
                 "capex_dep": 0.8 + 1.2 * ppe / ppe_scale + rng.normal(0, 0.1),
                 "ncx_sales": 0.02 + 0.1 * ppe / ppe_scale, "s_cap": float(rng.uniform(0.5, 3)),
                 "junk": float(rng.uniform(0, 1))}
            v["capspend"] = 0.02 + 0.06 * ppe + 0.02 * v["s_cap"]   # related to PP&E, not a copy of it
            out[y][ind] = v
    return out


def _rows(token, year, sim):
    """Header row + data rows of one table (industry names in column 0), with the planted defects."""
    cols = HEADERS[token][_era(year)]
    if token == "dbtfund" and _era(year) == "transition" and year != 2018:
        cols = [c for c in cols if c[1] != "de_rep"]           # Damodaran reports Debt/EBITDA in 2018 only here
    names = [_name(ind, year) for ind in INDUSTRIES]
    values = [[sim[year][ind][k] for _, k in cols] for ind in INDUSTRIES]
    if token == SHIFTED[0] and year == SHIFTED[1]:
        a, b = SHIFTED[2], SHIFTED[3]                            # names a..b-1 carry the data of the row above
        values[a:b] = values[a - 1:b - 1]
    if token == SCRAMBLED[0] and year == SCRAMBLED[1]:
        k = [c[1] for c in cols].index("ni")
        perm = np.random.default_rng(year).permutation(len(values))
        ni = [values[p][k] for p in perm]
        for r, v in zip(values, ni):
            r[k] = v
    total = [sum(r[j] for r in values) if isinstance(values[0][j], (int, float)) else None
             for j in range(len(cols))]
    return (["Industry Name"] + [h for h, _ in cols],
            [[n] + r for n, r in zip(names, values)] + [[AGGREGATE] + total])


def _write(path, year, token, sim):
    header, data = _rows(token, year, sim)
    wb = openpyxl.Workbook()
    era = _era(year)
    if era == "modern":
        faq = wb.active; faq.title = "Variables & FAQ"
        faq.append(["End Game", "What this dataset is for"]); faq.append([]); faq.append(["Variable", "Explanation"])
        ws = wb.create_sheet("Industry Averages")
        if token == "wacc" and year == 2019:                    # the real wacc19 has the sheets the other way round
            wb.move_sheet(ws, offset=-1)
    else:
        ws = wb.active; ws.title = "Sheet1"
    if era == "classic":
        ws.append(["Data from Value Line and Compustat"]); ws.append([])
    else:
        ws.append(["Date updated:", datetime(year + 1, 1, 5)])
        for label in ["Created by:", "What is this data?", "Home Page:", "Data website:"][: (4 if year >= 2017 else 1)]:
            ws.append([label, "..."])
    if token == "wacc":
        ws.append(["To update this spreadsheet, enter the following"]); ws.append(["Long Term Treasury bond rate", None, None, 0.04])
    ws.append([])
    ws.append(header)
    for r in data:
        ws.append(r)
    if era == "classic" and year <= 2004:
        wb.create_sheet("Sheet2"); wb.create_sheet("Sheet3")
    wb.save(path)


def write_fake_damodaran(folder, years=range(1999, 2026), seed=0, undated_newest="divfund"):
    """Write fake files for every dataset and year into `folder`; returns the simulated values.

    The newest year of `undated_newest` is saved without year digits (dated from its stamp).
    """
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    years = list(years)
    sim = _simulate(years, seed)
    for token in HEADERS:
        for y in years:
            name = f"{token}.xlsx" if (token == undated_newest and y == years[-1] and _era(y) != "classic") \
                else f"{token}{y % 100:02d}.xlsx"
            _write(folder / name, y, token, sim)
    return sim
